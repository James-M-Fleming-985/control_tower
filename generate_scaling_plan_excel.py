"""Generate the Scaling Plan Excel workbook for the Causal Affect project."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from datetime import date, timedelta

wb = openpyxl.Workbook()

# ── Styles ──────────────────────────────────────────────────────────────────
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=12)
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
SUBHEADER_FONT = Font(name="Calibri", bold=True, size=11)
MILESTONE_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
MILESTONE_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
TRACK_FILLS = {
    "A": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),  # green
    "B": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),  # orange
    "C": PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid"),  # blue
    "D": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),  # yellow
    "E": PatternFill(start_color="E2D9F3", end_color="E2D9F3", fill_type="solid"),  # purple
    "F": PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),  # pink
}


def style_header_row(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def style_cell(ws, row, col, fill=None, font=None):
    cell = ws.cell(row=row, column=col)
    cell.alignment = WRAP
    cell.border = THIN_BORDER
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    return cell


# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 1: Executive Summary
# ═══════════════════════════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Executive Summary"
ws.sheet_properties.tabColor = "2F5496"

headers = ["Dimension", "Current Baseline", "Final Target", "Gap", "Timeline"]
ws.append(headers)
style_header_row(ws, 1, 5)

baselines = [
    ["Data Volume", "~14K records, ~55 variables, single PostgreSQL (~50MB)", "10TB across distributed storage", "~99.999% gap", "12 months"],
    ["Model Accuracy (Direction)", "Unmeasured (~50% theoretical baseline)", "90%", "~40% improvement", "12 months"],
    ["Exploitation Metrics Accuracy", "Just implemented viability scoring, unvalidated", "90%", "~40% improvement", "12 months"],
    ["Deployment Time (E2E)", "30-45 min build + manual GitHub/Stripe/analytics = several hours", "2 hours fully automated", "Automation gap", "12 months"],
    ["Build Error Rate", "Unmeasured, no systematic tracking", "<5%", "Tracking + improvement needed", "12 months"],
    ["Pricing Intelligence", "Fixed tiers ($0-$199.99), no geo-pricing", "Region/economy-specific dynamic pricing", "Not started", "12 months"],
    ["Revenue Dashboard", "Stripe in systems3 only, no cross-app", "Central dashboard, all apps, all revenue streams", "Not started", "12 months"],
    ["Self-Improvement", "Legacy retrain code (non-operational)", "Continuous learning, auto-tuning, feedback loops", "Not started", "12 months"],
    ["Deployed Apps", "1 (Causal Affect) + 1 (systems3-project-reporter)", "4-10 apps (per AMP)", "2-8 more apps", "12 months"],
]

for i, row in enumerate(baselines, start=2):
    for j, val in enumerate(row, start=1):
        style_cell(ws, i, j)
        ws.cell(row=i, column=j, value=val)

ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 48
ws.column_dimensions["C"].width = 44
ws.column_dimensions["D"].width = 26
ws.column_dimensions["E"].width = 14

# Decisions row
r = len(baselines) + 3
ws.cell(row=r, column=1, value="Key Decisions").font = Font(bold=True, size=13)
decisions = [
    "Timeline: 12 months (~8 milestones at ~6 weeks each)",
    "App scope: 4-10 apps (per AMP plan)",
    "Data sources: All — social/financial feeds, news/search trends, IoT/sensor",
    "Storage evolution: PostgreSQL → TimescaleDB → S3 cold storage + hot cache",
]
for i, d in enumerate(decisions):
    ws.cell(row=r + 1 + i, column=1, value=d)

# ═══════════════════════════════════════════════════════════════════════════════
# Sheet 2: Milestone Roadmap (all milestones in one table)
# ═══════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("Milestone Roadmap")
ws2.sheet_properties.tabColor = "4472C4"

road_headers = ["Milestone", "Weeks", "Model Accuracy", "Exploitation Accuracy", "Build Errors", "Deploy Time", "Data Scale", "Revenue/Pricing"]
ws2.append(road_headers)
style_header_row(ws2, 1, 8)

milestones_summary = [
    ["M0: Instrument & Measure",   "1-3",   "Baseline (~50%)", "Baseline (~45-50%)", "Baseline (~20-30%)", "Current (5+ hrs)",   "~50MB",  "Audit Stripe baseline"],
    ["M1: Foundation Scaling",     "4-9",   "60%",             "55%",                "25%",                "Measured",           "100GB",  "Revenue Dashboard v1"],
    ["M2: Intelligence Layer",     "10-15", "65%",             "60%",                "20%",                "4 hours",            "500GB",  "Cross-app revenue tracking"],
    ["M3: Self-Learning Foundation","16-21","70%",             "65%",                "15%",                "3 hours",            "1TB",    "Pricing Intelligence v1"],
    ["M4: Scale & Optimize",       "22-27", "75%",             "70%",                "10%",                "2.5 hours",          "3TB",    "Revenue Dashboard v2 + P&L"],
    ["M5: Autonomous Operations",  "28-33", "80%",             "75%",                "5%",                 "2 hours",            "7TB",    "Dynamic Pricing Engine"],
    ["M6: Intelligence Engine",    "34-39", "85%",             "80%",                "Self-optimizing",    "Self-optimizing",    "10TB",   "Portfolio Optimization"],
    ["M7: Market Leadership",      "40-48", "90%",             "90%",                "<5% sustained",      "< 2 hrs sustained",  "10TB+",  "Full Financial OS ($20K+ MRR)"],
]

for i, row in enumerate(milestones_summary, start=2):
    for j, val in enumerate(row, start=1):
        c = style_cell(ws2, i, j, fill=MILESTONE_FILL if j == 1 else None,
                       font=MILESTONE_FONT if j == 1 else None)
        ws2.cell(row=i, column=j, value=val)

for col_idx in range(1, 9):
    ws2.column_dimensions[get_column_letter(col_idx)].width = [30, 8, 18, 22, 18, 18, 14, 30][col_idx - 1]


# ═══════════════════════════════════════════════════════════════════════════════
# Sheets 3-10: Individual Milestone Detail Sheets
# ═══════════════════════════════════════════════════════════════════════════════

milestones_detail = [
    {
        "name": "M0 - Instrument & Measure",
        "weeks": "Weeks 1-3",
        "motto": "You can't improve what you can't measure",
        "tracks": [
            ("A", "Model Accuracy Baseline", [
                "Run prediction validation on all matured predictions: POST /predictions/validate",
                "Record baseline direction_accuracy, avg_error_pct, avg_change_error_pct",
                "Add automated weekly validation cron (APScheduler or Railway cron)",
                "Add accuracy trend chart to Prediction Accuracy tab",
            ], "Documented baseline accuracy (~50-55%)"),
            ("B", "Build Error Baseline", [
                "Add BuildMetrics tracking to build_feature.py (syntax/test/frontend/wiring errors)",
                "Store results in JSON artifact uploaded with GitHub Actions",
                "Parse artifact in a /builds/metrics endpoint",
            ], "Error rate measurement framework, initial baseline (~20-30%)"),
            ("C", "Exploitation Metrics Baseline", [
                "Tag all 103 BUILD recommendations with viability scores",
                "Manually validate 10-20 BUILD recommendations against real-world outcomes",
                "Create exploitation_validation table for outcome tracking",
            ], "Documented baseline viability accuracy (~45-50%)"),
            ("D", "Revenue Baseline", [
                "Audit Stripe dashboard for current subscriber count and MRR",
                "Document current revenue per app",
            ], "Revenue baseline document"),
        ],
        "verification": [
            "All 4 baselines documented with numbers",
            "Automated weekly validation running",
            "Build error tracking live in CI/CD",
        ],
    },
    {
        "name": "M1 - Foundation Scaling",
        "weeks": "Weeks 4-9",
        "motto": "Target: +5% accuracy (55→60%), data to ~100GB",
        "tracks": [
            ("A", "Model Accuracy → 60%", [
                "Add daily Wikipedia pageview ingestion (55 vars × 365 days = 20K new records/year)",
                "Add daily Reddit activity ingestion for top 30 subreddits",
                "Implement walk-forward validation (train on N months, predict N+1, slide)",
                "Add multi-lag Granger testing (1-12 month lags)",
            ], "More granular data + proper validation → reduces overfitting"),
            ("B", "Build Errors → 25%", [
                "Add pre-build template validation (lint templates before generation)",
                "Add automated import resolution (detect missing imports, auto-add)",
                "Add structured error reporting to GitHub Actions artifacts",
            ], "Catch common failures before they happen"),
            ("C", "Exploitation Accuracy → 55%", [
                "Backtest viability scores against 6-month historical data",
                "Calibrate viability scoring weights based on backtest results",
                "Add target_growth_actual field for outcome tracking",
            ], "Empirical calibration of scoring model"),
            ("D", "Data Infrastructure", [
                "Enable TimescaleDB extension on Railway PostgreSQL",
                "Convert time_series_data to hypertable with time-based partitioning",
                "Add automated data retention policy (raw: 2 years, aggregated: forever)",
                "Set up S3 bucket for cold storage archive",
                "Connection pool tuning: pool_size=20, max_overflow=10",
            ], "~100GB target"),
            ("E", "Revenue Dashboard v1", [
                "Create RevenueEvent model (app_id, event_type, amount, currency, region, timestamp)",
                "Create /revenue/dashboard endpoint aggregating across all Stripe accounts",
                "Wire Stripe webhook events for all deployed apps to central tracker",
                "Build basic Revenue tab (MRR chart, subscriber count, per-app breakdown)",
            ], "Central webhook receiver that all deployed apps POST to"),
        ],
        "verification": [
            "Walk-forward validation shows 60%+ direction accuracy",
            "Build error rate measured with baseline established",
            "TimescaleDB partitioning active, daily ingestion running",
            "Revenue dashboard shows real MRR data",
        ],
    },
    {
        "name": "M2 - Intelligence Layer",
        "weeks": "Weeks 10-15",
        "motto": "Target: 65% accuracy, 60% exploitation, build time → 4hrs, errors → 20%",
        "tracks": [
            ("A", "Model Accuracy → 65%", [
                "Add ensemble model: Granger causality + linear regression + ARIMA",
                "Implement confidence-weighted ensemble (weight by historical accuracy)",
                "Add feature engineering: rolling averages (7d/30d/90d), rate of change, volatility",
                "Add cross-validation signal confidence (Wikipedia + Reddit agreement score)",
            ], "Ensemble reduces individual model bias"),
            ("B", "Build Errors → 20%", [
                "Add AI-powered code review step (LLM validates before commit)",
                "Add integration test templates for common patterns (CRUD, auth, Stripe)",
                "Fix top 5 most common build failure patterns from M1 data",
            ], "Systematic failure pattern elimination"),
            ("C", "Exploitation Accuracy → 60%", [
                "Add Google Trends integration (pytrends with proxy rotation)",
                "Cross-reference viability scores with Google Trends search volume",
                "Add competition scoring from real arXiv citation counts",
            ], "Richer data inputs → more accurate viability"),
            ("D", "Build Automation → 4 hours", [
                "Automate GitHub repo creation via GitHub API",
                "Auto-generate Railway project via Railway API",
                "Auto-wire Stripe subscription (create product + price via API)",
                "Auto-inject GA4 measurement ID from analytics template",
            ], "One-click Build & Deploy from Exploitation Board"),
            ("E", "Data Scale → 500GB", [
                "Add FRED economic indicators (50+ variables: GDP, CPI, unemployment)",
                "Add Alpha Vantage intraday data for top 10 stocks (5-min intervals)",
                "Add news sentiment from GDELT API (daily event counts by theme)",
                "Set up ETL pipeline with batch scheduling (APScheduler)",
            ], "More diverse, higher-frequency data sources"),
        ],
        "verification": [
            "Ensemble model shows 65%+ direction accuracy on walk-forward test",
            "Build → deploy → live app achievable in 4 hours with automation",
            "Revenue dashboard tracking deployed app(s)",
        ],
    },
    {
        "name": "M3 - Self-Learning Foundation",
        "weeks": "Weeks 16-21",
        "motto": "Target: 70% accuracy, 65% exploitation, build time → 3hrs, errors → 15%",
        "tracks": [
            ("A", "Model Accuracy → 70%", [
                "Implement automatic model retraining when accuracy drops below threshold",
                "Add prediction feedback loop: validated predictions → retrain features",
                "Add anomaly detection (Z-score) to flag unusual signals",
                "Implement adaptive lag selection (dynamically adjust based on recent accuracy)",
            ], "System learns from its own prediction outcomes"),
            ("B", "Build Errors → 15%", [
                "Build error pattern classifier (import, syntax, config, runtime categories)",
                "Auto-fix common patterns (missing __init__.py, wrong imports, missing env vars)",
                "Add scaffold regression tests (test each template produces working code)",
            ], "Error taxonomy → targeted fixes → prevent recurrence"),
            ("C", "Exploitation Accuracy → 65%", [
                "Track actual outcomes of PURSUING recommendations",
                "Add time-to-market analysis (recommendation → live product duration)",
                "Calibrate opportunity_duration_months against actual windows",
            ], "Closed-loop validation against real deployment outcomes"),
            ("D", "Build Automation → 3 hours", [
                "Add automated domain configuration (subdomain allocation)",
                "Add automated SSL certificate provisioning",
                "Add automated database setup for deployed apps",
                "Add post-deploy health check + smoke test",
            ], "Remove remaining manual steps"),
            ("E", "Data Scale → 1TB", [
                "Add social media feeds (Twitter/X academic API for trend data)",
                "Add IoT starter: OpenWeather API for 100 cities (hourly)",
                "Add Pushshift/Pullpush for full Reddit historical data",
                "Implement data compression in TimescaleDB (compress chunks >7 days old)",
            ], "Breadth of sources + compression for storage efficiency"),
            ("F", "Pricing Intelligence v1", [
                "Add IP geolocation to dashboard (MaxMind GeoLite2 free database)",
                "Map regions to PPP tiers from World Bank API",
                "Create pricing_recommendation table (region, suggested_price, confidence)",
                "Add traffic-based demand estimation per region",
            ], "Traffic patterns + economic data → region-aware pricing"),
        ],
        "verification": [
            "Self-retraining loop running automatically on weekly schedule",
            "Accuracy at 70%+ with feedback loop active",
            "Pricing intelligence generating region-specific suggestions",
            "1TB data threshold reached",
        ],
    },
    {
        "name": "M4 - Scale & Optimize",
        "weeks": "Weeks 22-27",
        "motto": "Target: 75% accuracy, 70% exploitation, build time → 2.5hrs, errors → 10%",
        "tracks": [
            ("A", "Model Accuracy → 75%", [
                "Add LSTM/transformer-based time series model (PyTorch) as ensemble member",
                "Implement feature importance ranking (SHAP values) for interpretability",
                "Add regime detection (bull/bear/sideways affect model choice)",
                "Implement online learning (update model incrementally with new data)",
            ], "Deep learning for non-linear patterns + regime-awareness"),
            ("B", "Build Errors → 10%", [
                "Implement AI-powered scaffold improvement (analyze fails → update templates)",
                "Add end-to-end integration test suite for full build pipeline",
                "Add canary deployment (deploy to staging first, promote if healthy)",
            ], "Self-improving templates based on failure analysis"),
            ("C", "Exploitation Accuracy → 70%", [
                "Add competitive landscape analysis (how many similar apps exist?)",
                "Cross-reference with Product Hunt / Indie Hackers for market validation",
                "Add user engagement prediction model (predict DAU/MAU)",
            ], "Market intelligence beyond statistical signals"),
            ("D", "Build Automation → 2.5 hours", [
                "Add automated testing pipeline for deployed apps",
                "Add automated monitoring setup (uptime, error rate, performance)",
                "Add automated changelog and release notes generation",
            ], "Reduce post-deployment manual work"),
            ("E", "Data Scale → 3TB", [
                "Add Bloomberg/Reuters alternative data feeds (if budget allows)",
                "Add satellite imagery analysis for economic indicators (nighttime lights)",
                "Implement data lakehouse pattern (S3 + DuckDB + PostgreSQL hot data)",
            ], "Alternative data + hybrid storage architecture"),
            ("F", "Revenue Dashboard v2", [
                "Add Google AdSense integration for ad-supported apps",
                "Add per-app P&L tracking (revenue - infrastructure cost)",
                "Add subscriber cohort analysis (retention by signup month)",
                "Add revenue forecasting from growth trends",
            ], "Full financial visibility across portfolio"),
        ],
        "verification": [
            "Deep learning ensemble at 75%+ accuracy",
            "Build → deploy consistently under 2.5 hours",
            "Revenue dashboard showing all revenue streams (subscriptions + ads)",
            "3TB data with hybrid hot/cold storage",
        ],
    },
    {
        "name": "M5 - Autonomous Operations",
        "weeks": "Weeks 28-33",
        "motto": "Target: 80% accuracy, 75% exploitation, build time → 2hrs, errors → 5%",
        "tracks": [
            ("A", "Model Accuracy → 80%", [
                "Implement multi-horizon forecasting (1 week, 1 month, 3 month predictions)",
                "Add cross-market contagion detection (one market shock propagating)",
                "Implement model selection autopilot (best model per signal-target pair)",
                "Add uncertainty quantification (prediction intervals, not just point estimates)",
            ], "Multi-horizon + automatic model selection + uncertainty"),
            ("B", "Build Errors → 5%", [
                "Implement self-healing builds (detect error → auto-apply fix → retry)",
                "Add template versioning with A/B testing (new vs old templates)",
                "Implement build quality gate (block deploy if quality < threshold)",
            ], "Self-healing + quality gates"),
            ("C", "Exploitation Accuracy → 75%", [
                "Add market timing model (when to launch, not just what to build)",
                "Add user acquisition cost prediction from similar product data",
                "Add churn prediction for deployed products",
            ], "Full lifecycle prediction (build + launch + grow + retain)"),
            ("D", "Build Automation → 2 hours", [
                "Full zero-touch pipeline: recommendation → build → deploy → monitor → scale",
                "Add automated A/B testing for deployed apps",
                "Add automated pricing optimization based on conversion data",
            ], "Complete automation of all manual steps"),
            ("E", "Data Scale → 7TB", [
                "Add real-time streaming (Kafka/Redis Streams for sub-minute data)",
                "Add sensor/IoT data ingestion (via MQTT broker)",
                "Implement federated querying across PostgreSQL + S3 + streaming",
            ], "Real-time data foundation for time-critical signals"),
            ("F", "Pricing Intelligence v2", [
                "Implement dynamic pricing engine (A/B test prices per region)",
                "Add currency conversion + PPP-adjusted pricing per country",
                "Add price elasticity estimation from traffic → conversion data",
                "Revenue-optimal pricing recommendations per app per region",
            ], "Data-driven pricing optimization"),
        ],
        "verification": [
            "80%+ accuracy with uncertainty intervals",
            "Zero-touch build-to-deploy in 2 hours",
            "Self-healing builds at <5% error rate",
            "Revenue dashboard with dynamic pricing active",
        ],
    },
    {
        "name": "M6 - Intelligence Engine",
        "weeks": "Weeks 34-39",
        "motto": "Target: 85% accuracy, 80% exploitation, self-optimizing",
        "tracks": [
            ("A", "Model Accuracy → 85%", [
                "Add reinforcement learning for trading strategy optimization",
                "Implement model distillation (compress ensemble into fast inference)",
                "Add causal discovery (beyond Granger to structural causal models)",
            ], "Advanced ML + causal inference"),
            ("B", "Self-Optimizing Builds", [
                "AI generates feature requirements from opportunity analysis (no YAML needed)",
                "Auto-select best scaffold and tech stack per opportunity",
                "Builds generate their own tests and validation suites",
            ], "AI writes its own specifications"),
            ("C", "Exploitation Accuracy → 80%", [
                "Portfolio optimization across all deployed apps",
                "Resource re-allocation recommendations (scale winners, sunset losers)",
                "Auto-sunset underperforming products",
            ], "Portfolio management approach"),
            ("D", "Data Scale → 10TB", [
                "Full multi-source streaming ingestion operational",
                "Automated data quality monitoring and anomaly detection",
                "Data lineage tracking (know provenance of every prediction)",
            ], "Enterprise-grade data platform"),
            ("E", "Revenue Engine", [
                "Central dashboard: portfolio MRR, CAC, LTV, churn by app",
                "Automated financial reporting (monthly P&L per product)",
                "Revenue alerts and anomaly detection",
                "Investment-ready metrics dashboard",
            ], "Full financial operating system"),
        ],
        "verification": [
            "85%+ accuracy on 3-month forward predictions",
            "Self-generating build specifications",
            "Portfolio-level revenue optimization active",
            "10TB data platform operational",
        ],
    },
    {
        "name": "M7 - Market Leadership",
        "weeks": "Weeks 40-48",
        "motto": "Target: 90% accuracy, 90% exploitation, market-ready",
        "tracks": [
            ("A", "Model Accuracy → 90%", [
                "Continuous learning pipeline (models retrain and auto-promote)",
                "Model marketplace (share/sell prediction models to other users)",
                "Multi-asset class prediction (stocks, crypto, commodities, real estate)",
            ], "Continuous improvement + market breadth"),
            ("B", "Full Automation", [
                "End-to-end: signal detected → app built → deployed → monetized → optimized",
                "Average deployment time consistently under 2 hours",
                "Build error rate sustained below 5%",
            ], "Mature, battle-tested automation"),
            ("C", "Exploitation Accuracy → 90%", [
                "Multi-factor scoring validated against 12+ months of outcomes",
                "Automated opportunity discovery (system finds BUILD opportunities)",
                "Cross-opportunity synergy detection (complementary products)",
            ], "Validated scoring + autonomous discovery"),
            ("D", "Revenue Targets", [
                "4-10 deployed apps generating subscription revenue",
                "Region-specific pricing active in 10+ markets",
                "Central revenue dashboard with real-time MRR tracking",
                "Target: $20K+ MRR across portfolio",
            ], "Portfolio scale + pricing optimization"),
        ],
        "verification": [
            "90% direction accuracy on walk-forward (6-month window)",
            "90% exploitation recommendation viability (validated)",
            "<5% build errors sustained over 3-month period",
            "2-hour average deployment time sustained",
            "Revenue dashboard live with all apps reporting",
            "Dynamic pricing active in multiple regions",
        ],
    },
]

for ms in milestones_detail:
    sheet_name = ms["name"][:31]  # Excel limit
    ws_ms = wb.create_sheet(sheet_name)
    ws_ms.sheet_properties.tabColor = "4472C4"

    # Title
    ws_ms.merge_cells("A1:E1")
    title_cell = ws_ms.cell(row=1, column=1, value=f"{ms['name']}  |  {ms['weeks']}")
    title_cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=14)
    title_cell.fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Motto
    ws_ms.merge_cells("A2:E2")
    ws_ms.cell(row=2, column=1, value=ms["motto"]).font = Font(italic=True, size=11)

    # Headers
    row = 4
    for c_idx, h in enumerate(["Track", "Focus Area", "Actions", "Method / Rationale"], start=1):
        cell = ws_ms.cell(row=row, column=c_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    row = 5
    for track in ms["tracks"]:
        track_letter, focus, actions, method = track
        fill = TRACK_FILLS.get(track_letter)
        style_cell(ws_ms, row, 1, fill=fill)
        ws_ms.cell(row=row, column=1, value=f"Track {track_letter}")
        style_cell(ws_ms, row, 2, fill=fill, font=SUBHEADER_FONT)
        ws_ms.cell(row=row, column=2, value=focus)
        style_cell(ws_ms, row, 3, fill=fill)
        ws_ms.cell(row=row, column=3, value="\n".join(f"• {a}" for a in actions))
        style_cell(ws_ms, row, 4, fill=fill)
        ws_ms.cell(row=row, column=4, value=method)
        # Set row height based on content
        ws_ms.row_dimensions[row].height = max(30, 16 * len(actions))
        row += 1

    # Verification section
    row += 1
    ws_ms.merge_cells(f"A{row}:D{row}")
    v_title = ws_ms.cell(row=row, column=1, value="Verification Criteria")
    v_title.font = SUBHEADER_FONT
    v_title.fill = SUBHEADER_FILL
    v_title.border = THIN_BORDER
    row += 1
    for v in ms["verification"]:
        ws_ms.merge_cells(f"A{row}:D{row}")
        ws_ms.cell(row=row, column=1, value=f"✓ {v}").border = THIN_BORDER
        row += 1

    ws_ms.column_dimensions["A"].width = 12
    ws_ms.column_dimensions["B"].width = 28
    ws_ms.column_dimensions["C"].width = 65
    ws_ms.column_dimensions["D"].width = 44


# ═══════════════════════════════════════════════════════════════════════════════
# Sheet: Architecture Evolution
# ═══════════════════════════════════════════════════════════════════════════════
ws_arch = wb.create_sheet("Architecture Evolution")
ws_arch.sheet_properties.tabColor = "70AD47"

arch_headers = ["Milestone", "Storage", "Model Pipeline", "Build Pipeline", "Revenue System"]
ws_arch.append(arch_headers)
style_header_row(ws_arch, 1, 5)

arch_data = [
    ["M0", "PostgreSQL (14K records, ~50MB)", "Granger causality only (granger_v1)", "Manual build_feature.py + manual GitHub/Stripe (~5+ hrs)", "Manual Stripe dashboard checks"],
    ["M1", "+ TimescaleDB extension + daily ingestion → 100GB", "+ Walk-forward validation + multi-lag testing", "+ Error tracking + pre-build validation", "+ Central RevenueEvent model + webhook aggregation"],
    ["M2", "+ Connection pool tuning + query optimization → 500GB", "+ Ensemble (Granger + linear + ARIMA) + feature engineering", "+ Auto GitHub repo + auto Railway + auto Stripe + auto GA4 (~4 hrs)", "+ Per-app MRR chart + subscriber counts"],
    ["M3", "+ Data compression + S3 cold storage → 1TB", "+ Feedback loop + auto-retrain + anomaly detection", "+ Auto domain + auto SSL + auto DB + smoke tests (~3 hrs)", "+ Geo-pricing intelligence v1"],
    ["M4", "+ DuckDB analytics layer + data lakehouse → 3TB", "+ LSTM/transformer + regime detection + online learning", "+ Staging/canary + auto monitoring (~2.5 hrs)", "+ Google Ads integration + per-app P&L + cohort analysis"],
    ["M5", "+ Kafka streaming + federated queries → 7TB", "+ Multi-horizon + uncertainty quantification + autopilot", "+ Self-healing + quality gates (~2 hrs)", "+ Dynamic pricing engine + A/B pricing"],
    ["M6", "Full hybrid: TimescaleDB + S3 + Kafka → 10TB", "+ Causal discovery + reinforcement learning", "+ AI-generated specs + auto stack selection", "+ Portfolio optimization + auto-sunset"],
    ["M7", "10TB+ operational", "+ Continuous learning + multi-asset + model marketplace", "+ Zero-touch end-to-end pipeline", "+ Full financial OS + investment-ready dashboard"],
]

for i, row in enumerate(arch_data, start=2):
    for j, val in enumerate(row, start=1):
        style_cell(ws_arch, i, j)
        ws_arch.cell(row=i, column=j, value=val)

ws_arch.column_dimensions["A"].width = 10
ws_arch.column_dimensions["B"].width = 48
ws_arch.column_dimensions["C"].width = 48
ws_arch.column_dimensions["D"].width = 52
ws_arch.column_dimensions["E"].width = 48


# ═══════════════════════════════════════════════════════════════════════════════
# Sheet: Files to Create/Modify
# ═══════════════════════════════════════════════════════════════════════════════
ws_files = wb.create_sheet("Key Files")
ws_files.sheet_properties.tabColor = "ED7D31"

files_headers = ["File", "Purpose"]
ws_files.append(files_headers)
style_header_row(ws_files, 1, 2)

files_data = [
    ["revenue_models.py", "RevenueEvent, SubscriptionMetric, PricingRecommendation models"],
    ["revenue_router.py", "Revenue dashboard endpoints"],
    ["revenue_webhook.py", "Central Stripe webhook receiver for all apps"],
    ["pricing_engine.py", "Region-aware pricing intelligence"],
    ["model_ensemble.py", "Ensemble prediction service"],
    ["auto_retrain.py", "Self-learning feedback loop"],
    ["streaming_ingestion.py", "Kafka/Redis Streams ingestion"],
    ["build_error_tracker.py", "Systematic build error classification"],
    ["self_healing_builder.py", "Auto-fix build failures"],
]

for i, row in enumerate(files_data, start=2):
    for j, val in enumerate(row, start=1):
        style_cell(ws_files, i, j)
        ws_files.cell(row=i, column=j, value=val)

ws_files.column_dimensions["A"].width = 28
ws_files.column_dimensions["B"].width = 55


# ═══════════════════════════════════════════════════════════════════════════════
# Sheet: MS Project (flat WBS with durations, dates, zero-duration milestones)
# ═══════════════════════════════════════════════════════════════════════════════
ws_msp = wb.create_sheet("MS Project", 1)  # Insert as 2nd sheet (after Executive Summary)
ws_msp.sheet_properties.tabColor = "70AD47"

msp_headers = ["WBS", "Task Name", "Duration (days)", "Start", "Finish", "Outline Level"]
ws_msp.append(msp_headers)
style_header_row(ws_msp, 1, 6)

PROJECT_START = date(2026, 3, 9)
DATE_FMT = "YYYY-MM-DD"

# Map milestone week ranges to start/end dates (work weeks, Mon-Fri)
def week_to_date(week_num):
    """Convert 1-based week number to Monday date relative to project start."""
    return PROJECT_START + timedelta(weeks=week_num - 1)

MILESTONE_WEEKS = [
    (1, 3),    # M0
    (4, 9),    # M1
    (10, 15),  # M2
    (16, 21),  # M3
    (22, 27),  # M4
    (28, 33),  # M5
    (34, 39),  # M6
    (40, 48),  # M7
]

msp_row = 2
ms_idx = 0

for ms in milestones_detail:
    ms_idx += 1
    wk_start, wk_end = MILESTONE_WEEKS[ms_idx - 1]
    ms_start = week_to_date(wk_start)
    ms_end = week_to_date(wk_end) + timedelta(days=4)  # Friday of last week
    ms_duration = (wk_end - wk_start + 1) * 5  # work days

    # --- Milestone summary row ---
    wbs = f"{ms_idx}"
    for c, val in enumerate([wbs, ms["name"], ms_duration, ms_start, ms_end, 1], start=1):
        cell = style_cell(ws_msp, msp_row, c, fill=MILESTONE_FILL, font=MILESTONE_FONT)
        ws_msp.cell(row=msp_row, column=c, value=val)
        if c in (4, 5):
            ws_msp.cell(row=msp_row, column=c).number_format = DATE_FMT
    msp_row += 1

    # --- Tracks & actions ---
    track_idx = 0
    for track in ms["tracks"]:
        track_idx += 1
        track_letter, focus, actions, method = track
        fill = TRACK_FILLS.get(track_letter)

        # Tracks run in parallel → same dates as milestone
        track_wbs = f"{ms_idx}.{track_idx}"
        for c, val in enumerate([track_wbs, f"Track {track_letter}: {focus}", ms_duration, ms_start, ms_end, 2], start=1):
            cell = style_cell(ws_msp, msp_row, c, fill=fill, font=SUBHEADER_FONT if c == 2 else None)
            ws_msp.cell(row=msp_row, column=c, value=val)
            if c in (4, 5):
                ws_msp.cell(row=msp_row, column=c).number_format = DATE_FMT
        msp_row += 1

        # Actions: divide milestone duration among actions sequentially
        num_actions = len(actions)
        action_dur = max(1, ms_duration // num_actions)
        action_start = ms_start
        for a_idx, action in enumerate(actions, start=1):
            a_wbs = f"{ms_idx}.{track_idx}.{a_idx}"
            # Last action gets remaining days to avoid rounding gaps
            if a_idx == num_actions:
                a_end = ms_end
                a_dur = max(1, (a_end - action_start).days + 1)
                # Approximate work days (exclude weekends)
                full_weeks = a_dur // 7
                remaining_days = a_dur % 7
                a_dur_work = full_weeks * 5 + min(remaining_days, 5)
            else:
                a_dur_work = action_dur
                a_end = action_start + timedelta(days=max(1, (a_dur_work * 7) // 5 - 1))

            for c, val in enumerate([a_wbs, action, a_dur_work, action_start, a_end, 3], start=1):
                cell = style_cell(ws_msp, msp_row, c)
                ws_msp.cell(row=msp_row, column=c, value=val)
                if c in (4, 5):
                    ws_msp.cell(row=msp_row, column=c).number_format = DATE_FMT
            msp_row += 1
            action_start = a_end + timedelta(days=1)

    # --- Milestone completion (zero-duration) ---
    ms_complete_wbs = f"{ms_idx}.{track_idx + 1}"
    for c, val in enumerate([ms_complete_wbs, f"{ms['name']} Complete", 0, ms_end, ms_end, 1], start=1):
        cell = style_cell(ws_msp, msp_row, c,
                          fill=PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid"),
                          font=Font(bold=True))
        ws_msp.cell(row=msp_row, column=c, value=val)
        if c in (4, 5):
            ws_msp.cell(row=msp_row, column=c).number_format = DATE_FMT
    msp_row += 1

ws_msp.column_dimensions["A"].width = 12
ws_msp.column_dimensions["B"].width = 65
ws_msp.column_dimensions["C"].width = 16
ws_msp.column_dimensions["D"].width = 14
ws_msp.column_dimensions["E"].width = 14
ws_msp.column_dimensions["F"].width = 14


# ═══════════════════════════════════════════════════════════════════════════════
# Save
# ═══════════════════════════════════════════════════════════════════════════════
output_path = "/workspaces/control_tower/cloned_repos/business_ventures/Scaling_Plan_Baseline_to_10TB_90pct.xlsx"
wb.save(output_path)
print(f"Saved: {output_path}")
