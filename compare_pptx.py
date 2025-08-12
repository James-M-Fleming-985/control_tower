from pptx import Presentation

# Load both presentations
new_pptx = Presentation('/workspaces/control_tower/cloned_repos/contract_projects/projects/Safran/powerpoint_reports/REACh_ZnNi_Line_Flash_Report_11082025_ControlTower.pptx')
old_pptx = Presentation('/workspaces/control_tower/cloned_repos/contract_projects/projects/Safran/powerpoint_reports/REACh_ZnNi_Line_Flash_Report_08082025_ControlTower.pptx')

# Helper function to find milestone tables in a presentation
def find_milestone_tables(prs):
    milestone_tables = []
    
    for slide_idx, slide in enumerate(prs.slides):
        for shape_idx, shape in enumerate(slide.shapes):
            if hasattr(shape, 'has_table') and shape.has_table:
                table = shape.table
                if len(table.rows) > 0 and len(table.columns) > 0:
                    header_text = table.cell(0,0).text_frame.text
                    if "milestone" in header_text.lower():
                        milestone_tables.append({
                            'slide_idx': slide_idx,
                            'shape_idx': shape_idx,
                            'table': table,
                            'header': header_text
                        })
    
    return milestone_tables

# Get milestone tables
new_milestone_tables = find_milestone_tables(new_pptx)
old_milestone_tables = find_milestone_tables(old_pptx)

print(f"Found {len(new_milestone_tables)} milestone tables in new presentation")
print(f"Found {len(old_milestone_tables)} milestone tables in old presentation")

# Compare milestone text between old and new presentations
for i, new_table_info in enumerate(new_milestone_tables):
    print(f"\n=========== NEW TABLE {i+1}: {new_table_info['header']} ===========")
    new_table = new_table_info['table']
    
    # Print column widths
    print(f"Column widths: ", end="")
    for col_idx, col in enumerate(new_table.columns):
        print(f"Col {col_idx}: {col.width}, ", end="")
    print()
    
    # Print milestone text in first column (typically where milestone names are)
    print("\nMilestone text (new presentation):")
    for row_idx, row in enumerate(new_table.rows):
        if row_idx > 1:  # Skip header rows
            cell = row.cells[0]
            text = cell.text_frame.text.strip()
            if text and text != "1":  # Skip empty cells
                # Check text frame config
                tf = cell.text_frame
                wrap = tf.word_wrap
                auto_size = tf.auto_size if hasattr(tf, 'auto_size') else 'Unknown'
                print(f"Row {row_idx}: {len(text)} chars - Wrap: {wrap}, Auto-size: {auto_size}")
                print(f"  TEXT: {text}")

# Check if we have matching tables in the old presentation
if i < len(old_milestone_tables):
    old_table_info = old_milestone_tables[i]
    old_table = old_table_info['table']
    
    print(f"\n----------- OLD TABLE {i+1}: {old_table_info['header']} -----------")
    
    # Print milestone text from old presentation
    print("\nMilestone text (old presentation):")
    for row_idx, row in enumerate(old_table.rows):
        if row_idx > 1:  # Skip header rows
            cell = row.cells[0]
            text = cell.text_frame.text.strip()
            if text and text != "1":  # Skip empty cells
                # Check text frame config
                tf = cell.text_frame
                wrap = tf.word_wrap
                auto_size = tf.auto_size if hasattr(tf, 'auto_size') else 'Unknown'
                print(f"Row {row_idx}: {len(text)} chars - Wrap: {wrap}, Auto-size: {auto_size}")
                print(f"  TEXT: {text}")
