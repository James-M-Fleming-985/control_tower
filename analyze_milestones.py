from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

def analyze_presentation(filepath):
    """Analyze the PowerPoint presentation to check milestone text display"""
    prs = Presentation(filepath)
    
    # Find milestone slides - looking for slides with tables
    for i, slide in enumerate(prs.slides):
        print(f"\nChecking Slide {i+1}:")
        tables_found = 0
        
        for j, shape in enumerate(slide.shapes):
            if hasattr(shape, 'has_table') and shape.has_table:
                table = shape.table
                tables_found += 1
                
                # Check if this could be a milestone table
                if len(table.rows) >= 3 and len(table.columns) >= 3:
                    first_cell_text = table.cell(0, 0).text_frame.text
                    print(f"  Table {tables_found}: '{first_cell_text}'")
                    
                    if "Milestone" in first_cell_text:
                        print(f"  MILESTONE TABLE FOUND!")
                        print(f"  Column count: {len(table.columns)}")
                        print(f"  Row count: {len(table.rows)}")
                        print(f"  Column widths: {[col.width for col in table.columns]}")
                        
                        # Print row heights
                        print(f"  Row heights: {[row.height for row in table.rows]}")
                        
                        # Analyze data cells (rows 2-4)
                        for row_idx in range(2, min(5, len(table.rows))):
                            # Get milestone text cell (first column)
                            cell = table.cell(row_idx, 0)
                            text = cell.text_frame.text.strip()
                            
                            if text and text != "—":  # Skip empty cells
                                text_len = len(text)
                                tf = cell.text_frame
                                
                                # Get text frame properties
                                auto_size = str(tf.auto_size) if hasattr(tf, 'auto_size') else 'Unknown'
                                word_wrap = str(tf.word_wrap) if hasattr(tf, 'word_wrap') else 'Unknown'
                                
                                print(f"\n  Row {row_idx+1} milestone: [{text_len} chars] '{text}'")
                                print(f"    Word wrap: {word_wrap}, Auto-size: {auto_size}")
                                
                                # For a non-empty milestone, check the first paragraph font properties
                                if hasattr(tf, 'paragraphs') and len(tf.paragraphs) > 0:
                                    p = tf.paragraphs[0]
                                    if hasattr(p, 'font') and hasattr(p.font, 'size'):
                                        font_size = p.font.size.pt if hasattr(p.font.size, 'pt') else 'Unknown'
                                        print(f"    Font size: {font_size}")
                                
                                # Report if text appears to be truncated (ends with '...')
                                if text.endswith('...'):
                                    print(f"    ⚠️ WARNING: Text appears to be truncated!")
        
        if tables_found == 0:
            print(f"  No tables found on this slide")

# Analyze the presentation
print("Analyzing new presentation...")
analyze_presentation('/workspaces/control_tower/cloned_repos/contract_projects/projects/Safran/powerpoint_reports/REACh_ZnNi_Line_Flash_Report_11082025_ControlTower.pptx')

print("\nAnalyzing previous presentation...")
try:
    analyze_presentation('/workspaces/control_tower/cloned_repos/contract_projects/projects/Safran/powerpoint_reports/REACh_ZnNi_Line_Flash_Report_08082025_ControlTower.pptx')
except Exception as e:
    print(f"Could not analyze previous presentation: {e}")
