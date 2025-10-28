#!/usr/bin/env python3
from openpyxl import load_workbook
from pathlib import Path

src = Path('/workspaces/control_tower/cloned_repos/professional_excellence/projects/PROJECT-001 NADCAP COMPLIANCE/inputs/NADCAP 2026 Action Plan-281025.xlsx')
if not src.exists():
    raise SystemExit(f"File not found: {src}")

wb = load_workbook(filename=str(src), data_only=False)
if 'Action Plan' not in wb.sheetnames or 'Gap Analysis' not in wb.sheetnames:
    raise SystemExit('Expected sheets "Action Plan" and "Gap Analysis" not both present')

ap = wb['Action Plan']
ga = wb['Gap Analysis']

# helper to find header index by substring
def find_col_by_substr(sheet, substrs):
    first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    for i, h in enumerate(first, start=1):
        if h is None:
            continue
        key = str(h).lower()
        for s in substrs:
            if s in key:
                return i
    return None

# Find columns
ap_clause_col = find_col_by_substr(ap, ['clause']) or 3
# prefer exact/documentation header for documentation reference; fall back to column 4 if detection is ambiguous
ap_docref_col = find_col_by_substr(ap, ['documentation reference'])
if ap_docref_col is None:
    # try looser match but avoid matching 'description' or 'action description'
    for i, h in enumerate(next(ap.iter_rows(min_row=1, max_row=1, values_only=True)), start=1):
        if h is None:
            continue
        key = str(h).lower()
        if 'documentation' in key or (('reference' in key or 'doc' in key) and 'clause' not in key):
            ap_docref_col = i
            break
if ap_docref_col is None:
    ap_docref_col = 4

ga_clause_col = find_col_by_substr(ga, ['clause']) or 6
ga_primary_ref_col = find_col_by_substr(ga, ['primary evidence doc ref']) or None
ga_primary_title_col = find_col_by_substr(ga, ['primary evidence title']) or None
if ga_primary_ref_col is None or ga_primary_title_col is None:
    # fallback by trying looser matches and defaulting to columns 10 and 11
    for i, h in enumerate(next(ga.iter_rows(min_row=1, max_row=1, values_only=True)), start=1):
        if h is None:
            continue
        key = str(h).lower()
        if 'primary evidence doc' in key or 'primary evidence doc ref' in key or ("doc ref" in key):
            ga_primary_ref_col = i
        if 'primary evidence title' in key or ('primary evidence' in key and 'title' in key):
            ga_primary_title_col = i
    if ga_primary_ref_col is None:
        ga_primary_ref_col = 10
    if ga_primary_title_col is None:
        ga_primary_title_col = 11

print('Action Plan clause col:', ap_clause_col, 'docref col:', ap_docref_col)
print('Gap Analysis clause col:', ga_clause_col, 'primary ref col:', ga_primary_ref_col, 'primary title col:', ga_primary_title_col)

# Build formula using INDEX/MATCH on the clause field in Action Plan (ap_clause_col)
# Assumption: Match on exact clause text. If multiple gap rows share same clause, first match used.

ap_max_row = ap.max_row
for r in range(2, ap_max_row+1):
    ap_clause_cell = ap.cell(row=r, column=ap_clause_col)
    # reference to this clause cell in formula (e.g., C2)
    clause_ref = ap_clause_cell.coordinate
    # construct formula: IFERROR(INDEX('Gap Analysis'!$J:$J, MATCH(TRIM(C2),'Gap Analysis'!$F:$F,0)) & " - " & INDEX('Gap Analysis'!$K:$K, MATCH(TRIM(C2),'Gap Analysis'!$F:$F,0)), "")
    ga_ref_col_letter = ga.cell(row=1, column=ga_primary_ref_col).column_letter
    ga_title_col_letter = ga.cell(row=1, column=ga_primary_title_col).column_letter
    ga_clause_col_letter = ga.cell(row=1, column=ga_clause_col).column_letter
    # Build formula that finds the row in Gap Analysis where the clause matches the Action Plan clause
    # and concatenates Primary Evidence Doc Ref and Primary Evidence Title with a separator.
    # Example formula produced for row 2:
    # =IFERROR(INDEX('Gap Analysis'!$J:$J, MATCH(TRIM(C2), 'Gap Analysis'!$F:$F, 0)) & " - " & INDEX('Gap Analysis'!$K:$K, MATCH(TRIM(C2), 'Gap Analysis'!$F:$F, 0)), "")
    formula = (
        f"=IFERROR(INDEX('Gap Analysis'!${ga_ref_col_letter}:${ga_ref_col_letter}, MATCH(TRIM({clause_ref}), 'Gap Analysis'!${ga_clause_col_letter}:${ga_clause_col_letter}, 0)) & \" - \" & "
        f"INDEX('Gap Analysis'!${ga_title_col_letter}:${ga_title_col_letter}, MATCH(TRIM({clause_ref}), 'Gap Analysis'!${ga_clause_col_letter}:${ga_clause_col_letter}, 0)), \"\")"
    )
    # assign formula to documentation reference cell
    ap.cell(row=r, column=ap_docref_col).value = formula

# Ask Excel to recalc on load
try:
    wb.calculation_properties.fullCalcOnLoad = True
except Exception:
    # older/newer openpyxl may not expose this; ignore
    pass

out = src.parent / (src.stem + '.formulas' + src.suffix)
wb.save(str(out))
print('Saved updated workbook with formulas to', out)
