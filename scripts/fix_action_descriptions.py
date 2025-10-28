#!/usr/bin/env python3
from openpyxl import load_workbook
from pathlib import Path

def find_col_by_substr(sheet, substrs):
    """Find column index by searching for substring in header row."""
    try:
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        for i, h in enumerate(first, start=1):
            if h is None:
                continue
            key = str(h).lower()
            for s in substrs:
                if s in key:
                    return i
    except StopIteration:
        pass
    return None

def main():
    src_path = Path('cloned_repos/professional_excellence/projects/PROJECT-001 NADCAP COMPLIANCE/inputs/NADCAP 2026 Action Plan-281025.xlsx')
    
    if not src_path.exists():
        raise SystemExit(f"File not found: {src_path}")

    wb = load_workbook(filename=str(src_path), data_only=False)
    
    if 'Action Plan' not in wb.sheetnames or 'Gap Analysis' not in wb.sheetnames:
        raise SystemExit('Expected sheets "Action Plan" and "Gap Analysis" not both present')

    ap = wb['Action Plan']
    ga = wb['Gap Analysis']

    # Find column indices
    ap_desc_col = find_col_by_substr(ap, ['action description', 'description']) or 2
    ap_clause_col = find_col_by_substr(ap, ['clause']) or 3

    ga_clause_col = find_col_by_substr(ga, ['clause']) or 6
    ga_primary_ref_col = find_col_by_substr(ga, ['primary evidence doc ref']) or 10
    ga_primary_title_col = find_col_by_substr(ga, ['primary evidence title']) or 11

    print(f'Action Plan: desc_col={ap_desc_col}, clause_col={ap_clause_col}')
    print(f'Gap Analysis: clause_col={ga_clause_col}, ref_col={ga_primary_ref_col}, title_col={ga_primary_title_col}')

    # Get column letters for formulas
    ga_ref_col_letter = ga.cell(row=1, column=ga_primary_ref_col).column_letter
    ga_title_col_letter = ga.cell(row=1, column=ga_primary_title_col).column_letter
    ga_clause_col_letter = ga.cell(row=1, column=ga_clause_col).column_letter

    # Process each row in Action Plan
    ap_max_row = ap.max_row
    for r in range(2, ap_max_row + 1):
        clause_cell = ap.cell(row=r, column=ap_clause_col)
        clause_ref = clause_cell.coordinate  # e.g., "C2"
        
        # Create formula for Action Description that pulls from Gap Analysis
        # Template: "Review document [REF] [TITLE] to ensure it adequately satisfies the requirements of NADCAP clause [CLAUSE]"
        formula = (
            f'="Review document " & IFERROR(INDEX(\'Gap Analysis\'!${ga_ref_col_letter}:${ga_ref_col_letter}, '
            f'MATCH(TRIM({clause_ref}), \'Gap Analysis\'!${ga_clause_col_letter}:${ga_clause_col_letter}, 0)), "") & '
            f'" " & IFERROR(INDEX(\'Gap Analysis\'!${ga_title_col_letter}:${ga_title_col_letter}, '
            f'MATCH(TRIM({clause_ref}), \'Gap Analysis\'!${ga_clause_col_letter}:${ga_clause_col_letter}, 0)), "") & '
            f'" to ensure it adequately satisfies the requirements of NADCAP clause " & {clause_ref}'
        )
        
        # Set the formula in Action Description column
        ap.cell(row=r, column=ap_desc_col).value = formula

    # Enable recalculation on load
    try:
        wb.calculation_properties.fullCalcOnLoad = True
    except Exception:
        pass

    # Save the updated workbook
    output_path = src_path.parent / (src_path.stem + '.fixed_descriptions' + src_path.suffix)
    wb.save(str(output_path))
    print(f'Saved updated workbook with Action Description formulas to: {output_path}')

if __name__ == '__main__':
    main()