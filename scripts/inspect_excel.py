#!/usr/bin/env python3
import sys
from openpyxl import load_workbook
from pathlib import Path

def find_header_cols(sheet):
    headers = {}
    first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    for idx, h in enumerate(first_row, start=1):
        if h is None:
            continue
        headers[h] = idx
    return headers


def guess_relevant_columns(headers):
    # Return header names that look like primary evidence or documentation refs
    matches = {}
    for h, idx in headers.items():
        key = h.lower()
        if any(k in key for k in ["primary", "evidence", "doc", "document", "reference", "title", "documentation"]):
            matches[h] = idx
    return matches


def print_sheet_info(wb, path):
    print(f"Workbook: {path}\nSheets: {wb.sheetnames}\n")
    for name in wb.sheetnames:
        sheet = wb[name]
        print(f"--- Sheet: '{name}' ---")
        try:
            headers = find_header_cols(sheet)
        except StopIteration:
            headers = {}
        if headers:
            print("Headers (first row):")
            for h, idx in headers.items():
                print(f"  {idx}: {h}")
        else:
            print("(no header row found)")

        relevant = guess_relevant_columns(headers)
        if relevant:
            print("Potential relevant columns:")
            for h, idx in relevant.items():
                print(f"  {idx}: {h}")
            # For each relevant column, print formulas and values in used range
            max_row = sheet.max_row
            for h, idx in relevant.items():
                col_letter = sheet.cell(row=1, column=idx).column_letter
                print(f"\nColumn {col_letter} ({h}) contents (showing formulas if present):")
                for r in range(2, min(max_row, 200)+1):
                    cell = sheet.cell(row=r, column=idx)
                    if cell.value is None:
                        continue
                    # openpyxl stores formula as a string beginning with '=' in cell.value when data_only=False
                    if isinstance(cell.value, str) and cell.value.startswith('='):
                        print(f"  {cell.coordinate}: FORMULA -> {cell.value}")
                    else:
                        print(f"  {cell.coordinate}: VALUE -> {cell.value}")
        else:
            print("No obvious 'primary evidence' or 'documentation reference' headers detected in first row.")
        # Also scan for formulas anywhere in sheet (few examples)
        print("\nScanning for formulas in sheet (first 50 found):")
        found = 0
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, max_col=sheet.max_column):
            for cell in row:
                if cell.value is None:
                    continue
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    print(f"  {cell.coordinate}: {cell.value}")
                    found += 1
                    if found >= 50:
                        break
            if found >= 50:
                break
        print("\n")

if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: inspect_excel.py <path-to-xlsx>")
        sys.exit(2)
    path = Path(sys.argv[1])
    if not path.exists():
        print(f"File not found: {path}")
        sys.exit(1)
    # load workbook with formulas (data_only=False) to see formula text
    wb = load_workbook(filename=str(path), data_only=False)
    print_sheet_info(wb, path)
    # Also load with data_only=True to show cached values
    wb_values = load_workbook(filename=str(path), data_only=True)
    print("Cached values (data_only=True) for first sheet, first 50 cells with values:")
    sheet = wb_values[wb_values.sheetnames[0]]
    count = 0
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_column, values_only=True):
        for v in row:
            if v is not None:
                print(v)
                count += 1
                if count >= 50:
                    break
        if count >= 50:
            break
