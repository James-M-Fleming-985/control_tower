#!/usr/bin/env python3
"""
PDF to Multiple Formats Converter
Converts NADCAP PDF to Word, Text, and HTML formats for better extraction
"""

import fitz  # PyMuPDF
import sys
import os
from pathlib import Path
import json
from docx import Document
import html

class PDFToMultipleFormats:
    def __init__(self, pdf_path):
        self.pdf_path = pdf_path
        self.doc = fitz.open(pdf_path)
        self.output_dir = Path("converted_formats")
        self.output_dir.mkdir(exist_ok=True)
        
    def convert_to_text(self):
        """Convert PDF to comprehensive text file with page markers"""
        print("📝 Converting to Text Format...")
        
        output_file = self.output_dir / "nadcap_complete.txt"
        
        with open(output_file, 'w', encoding='utf-8') as f:
            for page_num in range(len(self.doc)):
                page = self.doc[page_num]
                
                # Write page header
                f.write(f"\n{'='*80}\n")
                f.write(f"PAGE {page_num + 1}\n")
                f.write(f"{'='*80}\n\n")
                
                # Get text content
                text = page.get_text()
                
                # Clean and format text
                lines = text.split('\n')
                cleaned_lines = []
                
                for line in lines:
                    line = line.strip()
                    if line:  # Only keep non-empty lines
                        cleaned_lines.append(line)
                
                # Write cleaned content
                for line in cleaned_lines:
                    f.write(line + '\n')
                
                f.write('\n')  # Extra space between pages
        
        print(f"  ✅ Text file saved: {output_file}")
        return str(output_file)
    
    def convert_to_word(self):
        """Convert PDF to Word document with proper formatting"""
        print("📄 Converting to Word Format...")
        
        output_file = self.output_dir / "nadcap_complete.docx"
        
        doc = Document()
        
        # Add title
        title = doc.add_heading('NADCAP Audit Requirements - Complete Document', 0)
        
        for page_num in range(len(self.doc)):
            page = self.doc[page_num]
            
            # Add page heading
            doc.add_heading(f'Page {page_num + 1}', level=1)
            
            # Get text content
            text = page.get_text()
            
            # Process text in paragraphs
            paragraphs = text.split('\n\n')
            
            for paragraph in paragraphs:
                if paragraph.strip():
                    # Clean the paragraph
                    clean_para = ' '.join(paragraph.split())
                    if clean_para:
                        doc.add_paragraph(clean_para)
            
            # Add page break (except for last page)
            if page_num < len(self.doc) - 1:
                doc.add_page_break()
        
        doc.save(output_file)
        print(f"  ✅ Word file saved: {output_file}")
        return str(output_file)
    
    def convert_to_html(self):
        """Convert PDF to HTML with structure preservation"""
        print("🌐 Converting to HTML Format...")
        
        output_file = self.output_dir / "nadcap_complete.html"
        
        html_content = """
<!DOCTYPE html>
<html>
<head>
    <title>NADCAP Audit Requirements - Complete Document</title>
    <style>
        body { font-family: Arial, sans-serif; margin: 20px; }
        .page { border: 1px solid #ccc; margin: 20px 0; padding: 20px; }
        .page-header { background-color: #f0f0f0; padding: 10px; margin-bottom: 20px; }
        .clause-number { font-weight: bold; color: #0066cc; }
        .question { margin: 10px 0; }
        .guidance { font-style: italic; color: #666; }
        .yes-no { background-color: #ffffcc; padding: 5px; }
    </style>
</head>
<body>
    <h1>NADCAP Audit Requirements - Complete Document</h1>
"""
        
        for page_num in range(len(self.doc)):
            page = self.doc[page_num]
            
            html_content += f'\n    <div class="page">\n'
            html_content += f'        <div class="page-header"><h2>Page {page_num + 1}</h2></div>\n'
            
            # Get text content
            text = page.get_text()
            
            # Process lines
            lines = text.split('\n')
            
            for line in lines:
                line = line.strip()
                if line:
                    # Escape HTML
                    escaped_line = html.escape(line)
                    
                    # Identify and mark different types of content
                    if any(pattern in line for pattern in ['YES', 'NO', 'NA']):
                        html_content += f'        <div class="yes-no">{escaped_line}</div>\n'
                    elif line.startswith('Guidance:'):
                        html_content += f'        <div class="guidance">{escaped_line}</div>\n'
                    elif '?' in line:
                        html_content += f'        <div class="question">{escaped_line}</div>\n'
                    else:
                        # Check for clause numbers
                        import re
                        if re.search(r'\b\d+\.\d+(?:\.\d+)*\b', line):
                            html_content += f'        <div class="clause-number">{escaped_line}</div>\n'
                        else:
                            html_content += f'        <p>{escaped_line}</p>\n'
            
            html_content += '    </div>\n'
        
        html_content += '\n</body>\n</html>'
        
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(html_content)
        
        print(f"  ✅ HTML file saved: {output_file}")
        return str(output_file)
    
    def convert_to_json(self):
        """Convert PDF to structured JSON format"""
        print("📊 Converting to JSON Format...")
        
        output_file = self.output_dir / "nadcap_complete.json"
        
        document_data = {
            "title": "NADCAP Audit Requirements",
            "total_pages": len(self.doc),
            "pages": []
        }
        
        for page_num in range(len(self.doc)):
            page = self.doc[page_num]
            
            page_data = {
                "page_number": page_num + 1,
                "raw_text": page.get_text(),
                "lines": [],
                "potential_clauses": []
            }
            
            # Process lines
            text = page.get_text()
            lines = text.split('\n')
            
            for line in lines:
                line = line.strip()
                if line:
                    page_data["lines"].append(line)
                    
                    # Look for potential clause numbers
                    import re
                    clause_matches = re.findall(r'\b(\d+(?:\.\d+)*)\b', line)
                    for match in clause_matches:
                        if len(match.split('.')) >= 2:  # At least 2 parts (like 1.2)
                            page_data["potential_clauses"].append({
                                "clause_number": match,
                                "line": line
                            })
            
            document_data["pages"].append(page_data)
        
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump(document_data, f, indent=2, ensure_ascii=False)
        
        print(f"  ✅ JSON file saved: {output_file}")
        return str(output_file)
    
    def create_comprehensive_text_analysis(self):
        """Create a comprehensive text analysis file for manual review"""
        print("🔍 Creating Comprehensive Text Analysis...")
        
        output_file = self.output_dir / "nadcap_analysis.txt"
        
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write("NADCAP COMPREHENSIVE TEXT ANALYSIS\n")
            f.write("="*60 + "\n\n")
            
            all_potential_clauses = []
            
            for page_num in range(len(self.doc)):
                page = self.doc[page_num]
                text = page.get_text()
                
                f.write(f"\nPAGE {page_num + 1} ANALYSIS\n")
                f.write("-" * 30 + "\n")
                
                # Find all potential clause numbers
                import re
                clause_patterns = re.findall(r'\b(\d+(?:\.\d+)+)\b', text)
                unique_clauses = list(set(clause_patterns))
                unique_clauses.sort(key=lambda x: [int(part) for part in x.split('.')])
                
                if unique_clauses:
                    f.write(f"Potential clauses found: {len(unique_clauses)}\n")
                    for clause in unique_clauses:
                        f.write(f"  - {clause}\n")
                        all_potential_clauses.append((page_num + 1, clause))
                    f.write("\n")
                
                # Show lines with questions
                lines = text.split('\n')
                question_lines = [line.strip() for line in lines if '?' in line and line.strip()]
                
                if question_lines:
                    f.write(f"Questions found: {len(question_lines)}\n")
                    for i, question in enumerate(question_lines[:5]):  # First 5
                        f.write(f"  Q{i+1}: {question}\n")
                    if len(question_lines) > 5:
                        f.write(f"  ... and {len(question_lines) - 5} more\n")
                    f.write("\n")
                
                # Show YES/NO/NA lines
                yn_lines = [line.strip() for line in lines if any(word in line.upper() for word in ['YES', 'NO', 'NA']) and line.strip()]
                if yn_lines:
                    f.write(f"YES/NO/NA indicators: {len(yn_lines)}\n")
                
                f.write("\n" + "="*60 + "\n")
            
            # Summary
            f.write(f"\nOVERALL SUMMARY\n")
            f.write("="*30 + "\n")
            f.write(f"Total pages processed: {len(self.doc)}\n")
            f.write(f"Total potential clauses found: {len(all_potential_clauses)}\n")
            
            # Group by decimal levels
            decimal_levels = {}
            for page, clause in all_potential_clauses:
                level = len(clause.split('.'))
                if level not in decimal_levels:
                    decimal_levels[level] = []
                decimal_levels[level].append((page, clause))
            
            f.write("\nBreakdown by decimal levels:\n")
            for level in sorted(decimal_levels.keys()):
                f.write(f"  {level}-decimal clauses: {len(decimal_levels[level])}\n")
                examples = [clause for page, clause in decimal_levels[level][:5]]
                f.write(f"    Examples: {', '.join(examples)}\n")
        
        print(f"  ✅ Analysis file saved: {output_file}")
        return str(output_file)
    
    def convert_all_formats(self):
        """Convert PDF to all formats"""
        print("🚀 CONVERTING PDF TO MULTIPLE FORMATS")
        print("="*50)
        print(f"📄 Source: {self.pdf_path}")
        print(f"📁 Output directory: {self.output_dir}")
        print()
        
        results = {}
        
        try:
            results['text'] = self.convert_to_text()
            results['word'] = self.convert_to_word()
            results['html'] = self.convert_to_html()
            results['json'] = self.convert_to_json()
            results['analysis'] = self.create_comprehensive_text_analysis()
            
            print(f"\n🎯 CONVERSION COMPLETE!")
            print("="*30)
            print(f"📝 Text file: {results['text']}")
            print(f"📄 Word file: {results['word']}")
            print(f"🌐 HTML file: {results['html']}")
            print(f"📊 JSON file: {results['json']}")
            print(f"🔍 Analysis file: {results['analysis']}")
            
            print(f"\n💡 NEXT STEPS:")
            print("1. Review the text file for all clause numbers")
            print("2. Open the Word file for easy reading and editing")
            print("3. Use the HTML file for web-based review")
            print("4. Check the analysis file for comprehensive statistics")
            print("5. Use any of these formats for better extraction")
            
        except Exception as e:
            print(f"❌ Error during conversion: {e}")
            
        return results
    
    def close(self):
        """Close the PDF document"""
        self.doc.close()

def main():
    if len(sys.argv) != 2:
        print("Usage: python3 pdf_to_multiple_formats.py <path_to_pdf>")
        sys.exit(1)
    
    pdf_path = sys.argv[1]
    
    if not os.path.exists(pdf_path):
        print(f"Error: PDF file not found: {pdf_path}")
        sys.exit(1)
    
    converter = PDFToMultipleFormats(pdf_path)
    
    try:
        results = converter.convert_all_formats()
    finally:
        converter.close()

if __name__ == "__main__":
    main()
