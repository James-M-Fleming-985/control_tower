#!/usr/bin/env python3
"""
Surface Finishing Documentation Project - WBS PowerPoint Generator
Creates a professional WBS diagram in PowerPoint format
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR

def create_wbs_powerpoint():
    # Create presentation
    prs = Presentation()
    
    # Slide 1: Title Slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = title_slide.shapes.title
    subtitle = title_slide.placeholders[1]
    
    title.text = "Surface Finishing Documentation Project"
    subtitle.text = "Work Breakdown Structure (WBS)\nDeliverable-Oriented Hierarchy\nAugust - December 2025"
    
    # Slide 2: Complete WBS Hierarchy
    wbs_slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Add title
    title_shape = wbs_slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.8))
    title_frame = title_shape.text_frame
    title_frame.text = "SF Documentation Project - Work Breakdown Structure"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(24)
    title_para.font.bold = True
    title_para.alignment = PP_ALIGN.CENTER
    
    # WBS Level 1 - Project (Top)
    project_box = wbs_slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(3), Inches(1), Inches(3.5), Inches(0.8)
    )
    project_box.fill.solid()
    project_box.fill.fore_color.rgb = RGBColor(255, 235, 59)  # Yellow
    project_box.line.color.rgb = RGBColor(0, 0, 0)
    project_box.line.width = Pt(2)
    
    project_text = project_box.text_frame
    project_text.text = "1.0 SF Documentation\nStreamlining Project"
    project_text.paragraphs[0].font.size = Pt(12)
    project_text.paragraphs[0].font.bold = True
    project_text.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # WBS Level 2 - Major Deliverable Groups
    level2_data = [
        ("1.1 Project\nFoundation", 0.5),
        ("1.2 Gap Analysis\nSystem & Reports", 2.25),
        ("1.3 Documentation\nStructure", 4),
        ("1.4 Implementation\n& Rollout", 5.75),
        ("1.5 Closure &\nCompliance", 7.5)
    ]
    
    level2_boxes = []
    for text, x_pos in level2_data:
        box = wbs_slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x_pos), Inches(2.2), Inches(1.5), Inches(0.7)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(129, 199, 132)  # Green
        box.line.color.rgb = RGBColor(0, 0, 0)
        box.line.width = Pt(1)
        
        text_frame = box.text_frame
        text_frame.text = text
        text_frame.paragraphs[0].font.size = Pt(9)
        text_frame.paragraphs[0].font.bold = True
        text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        level2_boxes.append(box)
    
    # WBS Level 3 - Key Deliverable Packages (Selected)
    level3_data = [
        ("1.1.1 Project\nCharter Package", 0.2, 3.2),
        ("1.2.3 Gap Analysis\nSystem", 1.5, 3.2),
        ("1.2.4 Gap Reports\nPackage", 2.5, 3.2),
        ("1.3.1 Master Doc\nFramework", 3.7, 3.2),
        ("1.4.1 Final\nDocuments", 5.2, 3.2),
        ("1.4.2 System\nDeployment", 6.2, 3.2),
        ("1.5.1 Compliance\nAudit", 7.2, 3.2)
    ]
    
    for text, x_pos, y_pos in level3_data:
        box = wbs_slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x_pos), Inches(y_pos), Inches(1.2), Inches(0.6)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(100, 181, 246)  # Blue
        box.line.color.rgb = RGBColor(0, 0, 0)
        box.line.width = Pt(1)
        
        text_frame = box.text_frame
        text_frame.text = text
        text_frame.paragraphs[0].font.size = Pt(8)
        text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # WBS Level 4 - Specific Deliverables (Key ones)
    level4_data = [
        ("Approved\nProject Charter", 0.1, 4.1),
        ("Gap Analysis\nSoftware Tool", 1.4, 4.1),
        ("Excel Gap\nAnalysis Report", 2.4, 4.1),
        ("Overarching SF\nQMS Document", 3.6, 4.1),
        ("Final QMS\nDocument", 5.1, 4.1),
        ("Updated System\nStructure", 6.1, 4.1),
        ("AS9100 Compliance\nVerification", 7.1, 4.1)
    ]
    
    for text, x_pos, y_pos in level4_data:
        box = wbs_slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x_pos), Inches(y_pos), Inches(1), Inches(0.5)
        )
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(255, 171, 145)  # Orange
        box.line.color.rgb = RGBColor(0, 0, 0)
        box.line.width = Pt(1)
        
        text_frame = box.text_frame
        text_frame.text = text
        text_frame.paragraphs[0].font.size = Pt(7)
        text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Add connecting lines (simplified)
    # Project to Level 2
    for i, (_, x_pos) in enumerate(level2_data):
        line = wbs_slide.shapes.add_connector(
            1, Inches(4.75), Inches(1.8), Inches(x_pos + 0.75), Inches(2.2)
        )
        line.line.color.rgb = RGBColor(0, 0, 0)
        line.line.width = Pt(1)
    
    # Add timeline at bottom
    timeline_box = wbs_slide.shapes.add_textbox(Inches(0.5), Inches(5), Inches(9), Inches(0.8))
    timeline_frame = timeline_box.text_frame
    timeline_frame.text = "Timeline: Aug-Sep (Foundation) → Oct (Gap Analysis) → Nov (Documentation) → Dec (Implementation & Closure)"
    timeline_frame.paragraphs[0].font.size = Pt(11)
    timeline_frame.paragraphs[0].font.bold = True
    timeline_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Add legend
    legend_box = wbs_slide.shapes.add_textbox(Inches(0.5), Inches(5.8), Inches(9), Inches(1))
    legend_frame = legend_box.text_frame
    legend_frame.text = "Legend: 🟡 Project Level   🟢 Major Deliverable Groups   🔵 Deliverable Packages   🟠 Specific Deliverables"
    legend_frame.paragraphs[0].font.size = Pt(10)
    legend_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Slide 3: Critical Path & Resources
    critical_slide = prs.slides.add_slide(prs.slide_layouts[1])
    critical_title = critical_slide.shapes.title
    critical_title.text = "Critical Path & Resource Requirements"
    
    content = critical_slide.placeholders[1]
    content.text = """CRITICAL PATH DELIVERABLES:
• 1.2.3 Gap Analysis System - Foundation for automation
• 1.3.1 Master Documentation Framework - Core deliverable
• 1.4.1 Final Documents - Ready for deployment
• 1.5.1 Compliance Audit - Project validation

KEY RESOURCES NEEDED:
✓ Project Manager (James Fleming)
✓ Technical Author (Mike Warriner) 
✓ SME (James Bick)
🔴 Compliance Specialist - CRITICAL ADDITION NEEDED
• Documentation Software Specialist
• Training Coordinator

TIMELINE:
Aug-Sep: Foundation & Planning
Oct: Gap Analysis (Automated)
Nov: Documentation Development
Dec: Implementation & Closure"""
    
    # Slide 4: Success Metrics
    metrics_slide = prs.slides.add_slide(prs.slide_layouts[1])
    metrics_title = metrics_slide.shapes.title
    metrics_title.text = "Project Success Metrics & Deliverable Acceptance"
    
    metrics_content = metrics_slide.placeholders[1]
    metrics_content.text = """DELIVERABLE ACCEPTANCE CRITERIA:

📋 Project Charter Package
   → Signed by all stakeholders with scope/budget/timeline

🔧 Gap Analysis System
   → Functional tool comparing current vs AS9100 requirements

📊 Excel Gap Analysis Report
   → Complete spreadsheet with priorities and action plans

📖 Overarching SF QMS Document
   → AS9100-aligned structure with signpost framework

✅ AS9100 Compliance Verification
   → Independent audit confirming regulatory alignment

🎯 SUCCESS METRICS:
   • 100% AS9100 compliance verification
   • Automated gap analysis capability
   • Streamlined document hierarchy
   • Department training completion
   • Stakeholder approval & sign-off"""
    
    # Save presentation
    prs.save('/workspaces/control_tower/docs/SF_Documentation_WBS_Presentation.pptx')
    print("PowerPoint presentation created successfully!")
    print("File saved as: SF_Documentation_WBS_Presentation.pptx")

if __name__ == "__main__":
    create_wbs_powerpoint()
