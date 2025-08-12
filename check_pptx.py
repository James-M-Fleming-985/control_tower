from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR

prs = Presentation('/workspaces/control_tower/cloned_repos/contract_projects/projects/Safran/powerpoint_reports/REACh_ZnNi_Line_Flash_Report_11082025_ControlTower.pptx')
print(f'Total slides: {len(prs.slides)}')

# Find the milestone slide
# Analyze all slides with tables
milestone_slides = []
for i, slide in enumerate(prs.slides):
    has_table = False
    for shape in slide.shapes:
        if hasattr(shape, 'has_table') and shape.has_table:
            has_table = True
            table = shape.table
            print(f"Found table on Slide {i+1} with {len(table.rows)} rows and {len(table.columns)} columns")
            # Get first row, first cell text to help identify
            if len(table.rows) > 0 and len(table.columns) > 0:
                first_cell = table.cell(0,0).text_frame.text
                print(f"  First cell: '{first_cell}'")
                
                # Look for "milestone" in text or title
                if "milestone" in first_cell.lower() or "current" in first_cell.lower():
                    milestone_slides.append(i)
                    print(f"  This appears to be a milestone table!")
    
    if not has_table:
        print(f"No tables found on Slide {i+1}")

print(f"\nFound {len(milestone_slides)} potential milestone slides: {milestone_slides}")

# Set the milestone slide to analyze (use the first one found, or 5 as a fallback)
if milestone_slides:
    milestone_slide_idx = milestone_slides[0]
else:
    print("No clear milestone slides found, checking slide 5 as a fallback")
    milestone_slide_idx = 4  # 0-based index for slide 5

if milestone_slide_idx is not None:
    print(f"\nAnalyzing milestone slide {milestone_slide_idx+1}:")
    
    slide = prs.slides[milestone_slide_idx]
    for shape_idx, shape in enumerate(slide.shapes):
        if hasattr(shape, 'has_table') and shape.has_table:
            table = shape.table
            print(f"\nTable {shape_idx} with {len(table.rows)} rows and {len(table.columns)} columns")
            
            # Print column widths
            for col_idx, column in enumerate(table.columns):
                print(f"Column {col_idx} width: {column.width}")
            
            # Print table data with full milestone text
            header_row = True
            for row_idx, row in enumerate(table.rows):
                row_height = row.height
                print(f"\nRow {row_idx} (height: {row_height}):")
                
                for col_idx, cell in enumerate(row.cells):
                    text = cell.text_frame.text
                    text_length = len(text)
                    
                    # Get text frame properties
                    text_frame = cell.text_frame
                    wrap_info = "Wrap: " + str(text_frame.word_wrap)
                    auto_size_info = "Auto-size: " + str(text_frame.auto_size)
                    margin_info = f"Margins: L{text_frame.margin_left}/R{text_frame.margin_right}/T{text_frame.margin_top}/B{text_frame.margin_bottom}"
                    try:
                        anchor_info = f"Anchor: {text_frame.vertical_anchor}"
                    except:
                        anchor_info = "Anchor: Unknown"
                    
                    print(f"  Cell [{row_idx},{col_idx}] - {text_length} chars - {wrap_info}, {auto_size_info}, {margin_info}, {anchor_info}")
                    
                    # Show the full text for important milestone columns (typically column 1 or 2)
                    if col_idx in [0, 1] and row_idx > 0:  # Skip headers
                        if text_length > 40:
                            print(f"  FULL TEXT: {text}")
                
                header_row = False
else:
    print("No milestone table found")

# Also check the previous file
try:
    old_prs = Presentation('/workspaces/control_tower/cloned_repos/contract_projects/projects/Safran/powerpoint_reports/REACh_ZnNi_Line_Flash_Report_08082025_ControlTower.pptx')
    print("\n\nCOMPARISON WITH PREVIOUS FILE:")
    print(f"Previous file has {len(old_prs.slides)} slides")
except:
    print("\nCouldn't find previous file for comparison")
